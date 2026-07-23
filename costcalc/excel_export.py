"""
엑셀 출력 (openpyxl)
====================
계산 결과를 견적 검토 엑셀 서식으로 내보낸다.

시트 구성 (원가 도메인 검토 권고 반영):
  1. 견적요약   - 품목별 판가·매출(월/년)·적용율·프로젝트 정보
  2. 사출품원가 - 품목별 A/B/C/D 분해 (산출근거 포함)
  3. 조립품원가 - 품목별 A/B/C/D 분해 + 포장/운송 + 판가
  4. 판가시나리오 - 영업이익율 × 관리비율 판가 매트릭스
  5. 경쟁사비교 - 자사 vs 신흥정밀 양산판가 차이
  6. 손익BEP   - 고정/변동 분해, 공헌이익, BEP물량·매출
입력셀/계산셀 색 구분, 원화 서식, 합계행 강조, 각주(가정) 포함.
"""
from __future__ import annotations
import io
import re
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter

from .standards import MasterData
from .analysis import price_matrix, competitor_compare, bep_analysis, quote_summary

# --- 디자인 토큰 (UI와 동일 팔레트) -----------------------------------------
C_PRIMARY = "1F3A5F"     # 딥 네이비 (헤더)
C_ACCENT = "2E6B8A"      # 스틸 블루 (서브헤더)
C_TOTAL = "FBE9D0"       # 합계행 강조
C_INPUT = "FFF8E1"       # 입력셀
C_CALC = "F4F6F9"        # 계산셀
C_WHITE = "FFFFFF"
C_BORDER = "D0D7E2"

WON = '#,##0'
WON2 = '#,##0.00'
PCT = '0.0%'
PCT2 = '0.00%'

THIN = Side(style="thin", color=C_BORDER)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _title_font(size=11, color=C_WHITE, bold=True):
    return Font(name="맑은 고딕", size=size, bold=bold, color=color)


def _cell(ws, r, c, value, *, fill=None, font=None, fmt=None,
          align="left", bold=False, border=True):
    cell = ws.cell(row=r, column=c, value=value)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = font or Font(name="맑은 고딕", size=10, bold=bold,
                             color=C_WHITE if fill in (C_PRIMARY, C_ACCENT) else "1A1A1A")
    if fmt:
        cell.number_format = fmt
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=False)
    if border:
        cell.border = BORDER
    return cell


def _header_row(ws, r, headers, start=1, fill=C_PRIMARY):
    for i, h in enumerate(headers):
        _cell(ws, r, start + i, h, fill=fill, font=_title_font(), align="center")


def _section(ws, r, text, span, start=1, fill=C_ACCENT):
    _cell(ws, r, start, text, fill=fill, font=_title_font(), align="left")
    for c in range(start + 1, start + span):
        _cell(ws, r, c, "", fill=fill)
    ws.merge_cells(start_row=r, start_column=start, end_row=r, end_column=start + span - 1)


def _autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------------------------------------------------------------------------
# 1. 견적요약
# ---------------------------------------------------------------------------
def _sheet_summary(wb, result, master):
    ws = wb.active
    ws.title = "견적요약"
    ws.sheet_view.showGridLines = False
    proj = result.get("project", {})
    _autosize(ws, [22, 12, 10, 16, 30, 10, 10, 14, 14, 14, 18])

    _cell(ws, 1, 1, proj.get("title", "원가 견적 요약"),
          font=Font(name="맑은 고딕", size=15, bold=True, color=C_PRIMARY), border=False)
    ws.merge_cells("A1:K1")
    info = f'문서 {proj.get("doc_no","")}-{proj.get("revision","")}  |  고객사 {proj.get("customer","")}  |  작성일 {proj.get("date","")}  |  마스터 {master.config.get("version","")}'
    _cell(ws, 2, 1, info, font=Font(name="맑은 고딕", size=9, color="666666"), border=False)
    ws.merge_cells("A2:K2")

    summ = quote_summary(result["assemblies"])
    r = 4
    headers = ["품목", "고객", "Type", "자재코드", "품명", "관리비율", "영업이익율",
               "제조원가", "판가", "물량/월", "매출/월"]
    _header_row(ws, r, headers)
    r += 1
    for ln in summ["lines"]:
        _cell(ws, r, 1, ln["module"])
        _cell(ws, r, 2, ln["customer"], align="center")
        _cell(ws, r, 3, ln["type"], align="center")
        _cell(ws, r, 4, ln["material_code"], align="center")
        _cell(ws, r, 5, ln["name"])
        _cell(ws, r, 6, ln["admin_rate"], fmt=PCT, align="center")
        _cell(ws, r, 7, ln["profit_rate"], fmt=PCT, align="center")
        _cell(ws, r, 8, round(ln["mfg_cost"]), fmt=WON, align="right")
        _cell(ws, r, 9, round(ln["price"]), fmt=WON, align="right", bold=True)
        _cell(ws, r, 10, ln["volume"], fmt=WON, align="right")
        _cell(ws, r, 11, ln["sales_month"], fmt=WON, align="right")
        r += 1
    # 합계
    _cell(ws, r, 1, "합계", fill=C_TOTAL, bold=True)
    for c in range(2, 10):
        _cell(ws, r, c, "", fill=C_TOTAL)
    _cell(ws, r, 10, sum(l["volume"] for l in summ["lines"]), fill=C_TOTAL, fmt=WON, align="right", bold=True)
    _cell(ws, r, 11, summ["total_month"], fill=C_TOTAL, fmt=WON, align="right", bold=True)
    r += 2
    _cell(ws, r, 1, f'연간 매출 합계 : {summ["total_year"]:,.0f} 원',
          font=Font(name="맑은 고딕", size=11, bold=True, color=C_PRIMARY), border=False)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    r += 2
    _cell(ws, r, 1, "※ 판가 = 제조원가 × (1 + 관리비율 + 영업이익율) + 운반비",
          font=Font(name="맑은 고딕", size=9, italic=True, color="888888"), border=False)


# ---------------------------------------------------------------------------
# 제조원가 테이블 (사출/조립 공통 렌더)
# ---------------------------------------------------------------------------
def _cost_table(ws, res, master, is_assembly):
    ws.sheet_view.showGridLines = False
    _autosize(ws, [26, 12, 14, 14, 40])
    kind = "조립품" if is_assembly else "사출품"
    _cell(ws, 1, 1, f'[{kind}] {res["name"]}  제조원가 테이블',
          font=Font(name="맑은 고딕", size=13, bold=True, color=C_PRIMARY), border=False)
    ws.merge_cells("A1:E1")
    _cell(ws, 2, 1, f'자재코드 {res.get("material_code","")}  |  월물량 {res["volume"]:,}  |  일Capa {res["capa"]:,.1f}',
          font=Font(name="맑은 고딕", size=9, color="666666"), border=False)
    ws.merge_cells("A2:E2")
    r = 4

    # 재료비(A)
    _section(ws, r, "재료비 (A)", 5); r += 1
    _header_row(ws, r, ["항목", "재료/업체", "단가", "금액", "산출근거"], fill=C_ACCENT); r += 1
    for m in res["materials"]:
        _cell(ws, r, 1, m["name"], fill=C_INPUT)
        _cell(ws, r, 2, m.get("sub", ""), fill=C_INPUT, align="center")
        _cell(ws, r, 3, m["unit_price"], fill=C_INPUT, fmt=WON2, align="right")
        _cell(ws, r, 4, round(m["amount"], 2), fmt=WON2, align="right")
        _cell(ws, r, 5, m.get("basis", ""), font=Font(name="맑은 고딕", size=8, color="777777"))
        r += 1
    _cell(ws, r, 1, "재료비 합계(A)", fill=C_TOTAL, bold=True)
    _cell(ws, r, 2, "", fill=C_TOTAL); _cell(ws, r, 3, "", fill=C_TOTAL)
    _cell(ws, r, 4, round(res["material_total"], 2), fill=C_TOTAL, fmt=WON2, align="right", bold=True)
    _cell(ws, r, 5, "", fill=C_TOTAL); r += 2

    # 노무비(B)
    _section(ws, r, "노무비 (B)", 5); r += 1
    _header_row(ws, r, ["공정명", "투입인원", "", "금액", "산출근거"], fill=C_ACCENT); r += 1
    for p in res["labor"]["processes"]:
        _cell(ws, r, 1, p["name"], fill=C_INPUT)
        _cell(ws, r, 2, p["workers"], fill=C_INPUT, align="center")
        _cell(ws, r, 3, "", fill=C_CALC)
        _cell(ws, r, 4, round(p["amount"], 2), fmt=WON2, align="right")
        _cell(ws, r, 5, "임율 ÷ 일Capa × 인원", font=Font(name="맑은 고딕", size=8, color="777777"))
        r += 1
    _cell(ws, r, 1, "간접노무비", fill=C_CALC)
    _cell(ws, r, 2, "", fill=C_CALC); _cell(ws, r, 3, "", fill=C_CALC)
    _cell(ws, r, 4, round(res["labor"]["indirect"], 2), fmt=WON2, align="right")
    _cell(ws, r, 5, f'직접노무비 × {res["labor"]["indirect_rate"]:.0%}', font=Font(name="맑은 고딕", size=8, color="777777"))
    r += 1
    _cell(ws, r, 1, "노무비 합계(B)", fill=C_TOTAL, bold=True)
    _cell(ws, r, 2, "", fill=C_TOTAL); _cell(ws, r, 3, "", fill=C_TOTAL)
    _cell(ws, r, 4, round(res["labor"]["total"], 2), fill=C_TOTAL, fmt=WON2, align="right", bold=True)
    _cell(ws, r, 5, "", fill=C_TOTAL); r += 2

    # 제조경비(C)
    _section(ws, r, "제조경비 (C)", 5); r += 1
    _header_row(ws, r, ["항목", "", "", "금액", "산출근거"], fill=C_ACCENT); r += 1
    for o in res["overhead"]:
        _cell(ws, r, 1, o["name"], fill=C_CALC)
        _cell(ws, r, 2, "", fill=C_CALC); _cell(ws, r, 3, "", fill=C_CALC)
        _cell(ws, r, 4, round(o["amount"], 2), fmt=WON2, align="right")
        _cell(ws, r, 5, o.get("basis", ""), font=Font(name="맑은 고딕", size=8, color="777777"))
        r += 1
    _cell(ws, r, 1, "제조경비 합계(C)", fill=C_TOTAL, bold=True)
    _cell(ws, r, 2, "", fill=C_TOTAL); _cell(ws, r, 3, "", fill=C_TOTAL)
    _cell(ws, r, 4, round(res["overhead_total"], 2), fill=C_TOTAL, fmt=WON2, align="right", bold=True)
    _cell(ws, r, 5, "", fill=C_TOTAL); r += 2

    # 제조원가(D)
    _section(ws, r, "제조원가 (D) = A + B + C + LOSS", 5); r += 1
    for label, val, ratio in [
        ("재료비 합계(A)", res["material_total"], res["composition"]["material"]),
        ("노무비 합계(B)", res["labor"]["total"], res["composition"]["labor"]),
        ("제조경비 합계(C)", res["overhead_total"], res["composition"]["overhead"]),
        (f'LOSS ({res["loss_rate"]:.0%})', res["loss"], res["composition"]["loss"]),
    ]:
        _cell(ws, r, 1, label, fill=C_CALC)
        _cell(ws, r, 2, "", fill=C_CALC); _cell(ws, r, 3, "", fill=C_CALC)
        _cell(ws, r, 4, round(val, 2), fmt=WON2, align="right")
        _cell(ws, r, 5, ratio, fmt=PCT, align="center")
        r += 1
    _cell(ws, r, 1, "제조원가", fill=C_PRIMARY, font=_title_font(), bold=True)
    _cell(ws, r, 2, "", fill=C_PRIMARY); _cell(ws, r, 3, "", fill=C_PRIMARY)
    _cell(ws, r, 4, round(res["manufacturing_cost"], 2), fill=C_PRIMARY,
          font=_title_font(), fmt=WON2, align="right")
    _cell(ws, r, 5, "100%", fill=C_PRIMARY, font=_title_font(), align="center")
    r += 2

    # 조립품: 판가
    if is_assembly:
        q = res["quote"]
        _section(ws, r, "판가 산출", 5); r += 1
        for label, val, basis in [
            ("제조원가", res["manufacturing_cost"], ""),
            (f'관리비 ({q["admin_rate"]:.1%})', q["admin_cost"], "제조원가 × 관리비율"),
            (f'영업이익 ({q["profit_rate"]:.1%})', q["profit"], "제조원가 × 영업이익율"),
            ("운반비", q["transport"], "1회운임 ÷ (총적재 × 용적율)"),
        ]:
            _cell(ws, r, 1, label, fill=C_CALC)
            _cell(ws, r, 2, "", fill=C_CALC); _cell(ws, r, 3, "", fill=C_CALC)
            _cell(ws, r, 4, round(val, 2), fmt=WON2, align="right")
            _cell(ws, r, 5, basis, font=Font(name="맑은 고딕", size=8, color="777777"))
            r += 1
        _cell(ws, r, 1, "판  가", fill=C_TOTAL, bold=True)
        _cell(ws, r, 2, "", fill=C_TOTAL); _cell(ws, r, 3, "", fill=C_TOTAL)
        _cell(ws, r, 4, round(q["price"], 2), fill=C_TOTAL, fmt=WON2, align="right", bold=True)
        _cell(ws, r, 5, "", fill=C_TOTAL)


_FORBIDDEN = re.compile(r'[\\/?*\[\]:]')


def _safe_sheet_name(raw, used_names):
    """엑셀 시트명 규칙: 금지문자(\\ / ? * [ ] :) 제거, 31자 이내, 중복 회피."""
    name = _FORBIDDEN.sub("_", raw)[:31].strip() or "Sheet"
    base, i = name, 1
    while name.lower() in used_names:
        suffix = f"_{i}"
        name = base[:31 - len(suffix)] + suffix
        i += 1
    used_names.add(name.lower())
    return name


def _sheet_cost(wb, res, master, is_assembly, used_names):
    prefix = "조립_" if is_assembly else "사출_"
    title = _safe_sheet_name(prefix + res["name"], used_names)
    ws = wb.create_sheet(title)
    _cost_table(ws, res, master, is_assembly)


# ---------------------------------------------------------------------------
# 4. 판가시나리오
# ---------------------------------------------------------------------------
def _sheet_scenarios(wb, result, master):
    ws = wb.create_sheet("판가시나리오")
    ws.sheet_view.showGridLines = False
    _cell(ws, 1, 1, "판가 민감도 (영업이익율 × 관리비율)",
          font=Font(name="맑은 고딕", size=13, bold=True, color=C_PRIMARY), border=False)
    r = 3
    q = master.config["quote"]
    profits = q["scenario_profit_rates"]
    _autosize(ws, [26, 14, 14, 14, 14, 14, 14])
    for a in result["assemblies"]:
        _cell(ws, r, 1, f'{a["name"]}  (제조원가 {a["manufacturing_cost"]:,.0f})',
              fill=C_ACCENT, font=_title_font(), align="left")
        for c in range(2, 2 + len(profits)):
            _cell(ws, r, c, "", fill=C_ACCENT)
        r += 1
        _cell(ws, r, 1, "관리비율 \\ 영업이익율", fill=C_PRIMARY, font=_title_font(), align="center")
        for i, pr in enumerate(profits):
            _cell(ws, r, 2 + i, pr, fill=C_PRIMARY, font=_title_font(), fmt=PCT, align="center")
        r += 1
        mat = price_matrix(a["manufacturing_cost"], a["quote"]["transport"], master)
        for row in mat["rows"]:
            _cell(ws, r, 1, row["admin_rate"], fmt=PCT, align="center", fill=C_INPUT)
            for i, cell in enumerate(row["cells"]):
                is_base = (abs(cell["admin_rate"] - a["quote"]["admin_rate"]) < 1e-9 and
                           abs(cell["profit_rate"] - a["quote"]["profit_rate"]) < 1e-9)
                _cell(ws, r, 2 + i, round(cell["price"]), fmt=WON, align="right",
                      fill=C_TOTAL if is_base else None, bold=is_base)
            r += 1
        r += 1


# ---------------------------------------------------------------------------
# 5. 경쟁사비교
# ---------------------------------------------------------------------------
def _sheet_competitor(wb, result):
    ws = wb.create_sheet("경쟁사비교")
    ws.sheet_view.showGridLines = False
    _cell(ws, 1, 1, "경쟁사(신흥정밀) 양산판가 대비 비교",
          font=Font(name="맑은 고딕", size=13, bold=True, color=C_PRIMARY), border=False)
    _autosize(ws, [30, 12, 14, 16, 14, 14, 14, 12])
    r = 3
    _header_row(ws, r, ["품명", "고객", "자재코드", "당사 견적가", "경쟁사가", "차이", "차이%"]); r += 1
    for a in result["assemblies"]:
        for row in competitor_compare(a):
            _cell(ws, r, 1, row["name"])
            _cell(ws, r, 2, row["customer"], align="center")
            _cell(ws, r, 3, row["material_code"], align="center")
            _cell(ws, r, 4, round(row["our_price"]), fmt=WON, align="right", bold=True)
            _cell(ws, r, 5, row["competitor_price"], fmt=WON, align="right")
            _cell(ws, r, 6, round(row["diff"]) if row["diff"] is not None else "-",
                  fmt=WON, align="right")
            _cell(ws, r, 7, row["diff_pct"] if row["diff_pct"] is not None else "-",
                  fmt=PCT2, align="center")
            r += 1


# ---------------------------------------------------------------------------
# 6. 손익BEP
# ---------------------------------------------------------------------------
def _sheet_bep(wb, result, master):
    ws = wb.create_sheet("손익BEP")
    ws.sheet_view.showGridLines = False
    _cell(ws, 1, 1, "손익분기(BEP) / 공헌이익 분석",
          font=Font(name="맑은 고딕", size=13, bold=True, color=C_PRIMARY), border=False)
    _cell(ws, 2, 1, "※ 재료·복리후생·수도광열·소모품·포장·운반=변동비 / 감가·건물·수선·기타·관리비=고정비 (노무비 전액 변동 가정)",
          font=Font(name="맑은 고딕", size=8, italic=True, color="888888"), border=False)
    _autosize(ws, [30, 12, 14, 14, 14, 14, 14, 14])
    r = 4
    _header_row(ws, r, ["품목", "물량/월", "판가", "개당변동비", "개당공헌이익",
                        "공헌이익율", "BEP물량", "BEP매출"]); r += 1
    for a in result["assemblies"]:
        b = bep_analysis(a, master)
        _cell(ws, r, 1, a["name"])
        _cell(ws, r, 2, b["volume"], fmt=WON, align="right")
        _cell(ws, r, 3, round(b["price"]), fmt=WON, align="right")
        _cell(ws, r, 4, round(b["unit_variable"]), fmt=WON, align="right")
        _cell(ws, r, 5, round(b["contribution"]), fmt=WON, align="right", bold=True)
        _cell(ws, r, 6, b["cm_ratio"], fmt=PCT, align="center")
        _cell(ws, r, 7, b["bep_qty"] if b["bep_qty"] else "-", fmt=WON, align="right")
        _cell(ws, r, 8, b["bep_sales"] if b["bep_sales"] else "-", fmt=WON, align="right")
        r += 1


# ---------------------------------------------------------------------------
# 진입점
# ---------------------------------------------------------------------------
def build_workbook(result, master: MasterData) -> Workbook:
    wb = Workbook()
    _sheet_summary(wb, result, master)
    used = set()
    for p in result["parts"]:
        _sheet_cost(wb, p, master, False, used)
    for a in result["assemblies"]:
        _sheet_cost(wb, a, master, True, used)
    _sheet_scenarios(wb, result, master)
    _sheet_competitor(wb, result)
    _sheet_bep(wb, result, master)
    return wb


def export_bytes(result, master: MasterData) -> bytes:
    wb = build_workbook(result, master)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
